#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
report_md_to_docx.py — 把《SCI 热点解读报告》的 Markdown 稿转为 .docx

为什么单独写：每周报告是一份结构化长文（标题层级 + 大量表格 + 引用块 + 代码块）。
用通用转换链（pandoc / LibreOffice）在本机不可用，而报告的表格式 IF 分布、
方法与数据来源声明必须保持为**真表格**（不是截图、不是纯文本），
否则「可核验」这一条在 Word 侧就断了。

用法
  python3 report_md_to_docx.py --md 02_热点报告/xxx.md --out 02_热点报告/xxx.docx

支持：ATX 标题、GFM 表格、引用块、无序列表、有序列表、水平线、
      行内 **粗体** / *斜体* / `等宽`。其余语法按普通段落处理。
"""
import argparse
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

CN_FONT = "宋体"
CN_HEAD = "黑体"
MONO = "Consolas"

INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`)")


def set_cn(run, latin="Times New Roman", ea=CN_FONT):
    run.font.name = latin
    run._element.rPr.rFonts.set(qn("w:eastAsia"), ea)


def _add_runs(p, text, base_size=None, bold=False, italic=False):
    """按反引号切分写入段落：反引号内用等宽字体，其余用正文字体。

    修复：加粗/斜体区间内嵌行内代码（如 **`>5` 占 68.0%**）时，
    旧实现只剥掉外层 ** 或 *，内层反引号会原样落进 Word。
    """
    parts = text.split("`")
    for idx, seg in enumerate(parts):
        if seg == "":
            continue
        r = p.add_run(seg)
        if idx % 2 == 1:
            set_cn(r, latin=MONO, ea=MONO)
            r.font.size = Pt((base_size or 10.5) - 1)
        else:
            set_cn(r)
            if base_size:
                r.font.size = Pt(base_size)
        r.bold = bold
        r.italic = italic


def add_inline(p, text, base_size=None):
    """把含行内标记的文本写入段落（保留中英文混排）。"""
    for tok in INLINE.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            _add_runs(p, tok[2:-2], base_size, bold=True)
        elif tok.startswith("`") and tok.endswith("`") and len(tok) > 2:
            _add_runs(p, tok, base_size)
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            _add_runs(p, tok[1:-1], base_size, italic=True)
        else:
            _add_runs(p, tok, base_size)


def split_row(line):
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    return cells


def is_sep(line):
    return bool(re.match(r"^\|[\s:\-|]+\|\s*$", line.strip()))


def build(md_text, out_path, title=""):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)

    lines = md_text.split("\n")
    i, in_code = 0, False
    while i < len(lines):
        ln = lines[i]
        s = ln.strip()

        # 代码块
        if s.startswith("```"):
            in_code = not in_code
            i += 1
            continue
        if in_code:
            p = doc.add_paragraph()
            r = p.add_run(ln); set_cn(r, latin=MONO, ea=MONO); r.font.size = Pt(9)
            p.paragraph_format.left_indent = Pt(18)
            p.paragraph_format.space_after = Pt(0)
            i += 1
            continue

        # 表格
        if s.startswith("|") and i + 1 < len(lines) and is_sep(lines[i + 1]):
            header = split_row(lines[i])
            body, j = [], i + 2
            while j < len(lines) and lines[j].strip().startswith("|"):
                body.append(split_row(lines[j])); j += 1
            tb = doc.add_table(rows=1, cols=len(header))
            tb.style = "Table Grid"
            tb.alignment = WD_TABLE_ALIGNMENT.CENTER
            for k, h in enumerate(header):
                cell = tb.rows[0].cells[k]
                cell.text = ""
                add_inline(cell.paragraphs[0], h, base_size=9.5)
                for r in cell.paragraphs[0].runs:
                    r.bold = True
            for row in body:
                cells = tb.add_row().cells
                for k in range(len(header)):
                    cells[k].text = ""
                    add_inline(cells[k].paragraphs[0],
                               row[k] if k < len(row) else "", base_size=9.5)
            doc.add_paragraph()
            i = j
            continue

        # 标题
        m = re.match(r"^(#{1,6})\s+(.*)$", s)
        if m:
            lvl, txt = len(m.group(1)), m.group(2)
            h = doc.add_heading(level=min(lvl, 4))
            add_inline(h, txt)
            for r in h.runs:
                set_cn(r, ea=CN_HEAD)
            i += 1
            continue

        # 水平线
        if re.match(r"^-{3,}$", s):
            doc.add_paragraph()
            i += 1
            continue

        # 引用块
        if s.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            add_inline(p, s.lstrip(">").strip())
            for r in p.runs:
                r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
            i += 1
            continue

        # 列表
        if re.match(r"^[-*]\s+", s):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, re.sub(r"^[-*]\s+", "", s))
            i += 1
            continue
        if re.match(r"^\d+\.\s+", s):
            p = doc.add_paragraph(style="List Number")
            add_inline(p, re.sub(r"^\d+\.\s+", "", s))
            i += 1
            continue

        if not s:
            i += 1
            continue

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_inline(p, s)
        i += 1

    doc.save(out_path)
    return out_path


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--md", required=True)
    ap.add_argument("--out", required=True)
    a = ap.parse_args()
    md = open(a.md, encoding="utf-8").read()
    print("已生成:", build(md, a.out))
    return 0


if __name__ == "__main__":
    sys.exit(main())
